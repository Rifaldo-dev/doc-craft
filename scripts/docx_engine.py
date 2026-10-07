import os
import sys
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BLACK = RGBColor(0, 0, 0)

class StandardDocument:
    """
    StandardDocument Builder:
    Menghasilkan dokumen Microsoft Word (.docx) formal berstandar akademik dan industri.
    - 100% teks hitam murni (RGB 0,0,0)
    - Anti-slop (filter otomatis em-dash)
    - Tipografi konsisten (default Times New Roman)
    - Tabel bergaris bersih dan rapi
    - Manajemen margin presisi
    """
    def __init__(self, font_name="Times New Roman", margin_preset="academic"):
        self.doc = docx.Document()
        self.font_name = font_name
        self._setup_margins(margin_preset)
        self._setup_default_styles()

    def _setup_margins(self, preset):
        for section in self.doc.sections:
            if preset == "thesis_4_4_3_3": # Skripsi/Tesis: Kiri 4cm, Atas 4cm, Kanan 3cm, Bawah 3cm
                section.top_margin = Inches(1.57)
                section.bottom_margin = Inches(1.18)
                section.left_margin = Inches(1.57)
                section.right_margin = Inches(1.18)
            else: # Standar Akademik & Laporan: Kiri 3cm, Atas 2.5cm, Bawah 2.5cm, Kanan 2.5cm
                section.top_margin = Inches(1.0)
                section.bottom_margin = Inches(1.0)
                section.left_margin = Inches(1.18)
                section.right_margin = Inches(1.0)

    def _setup_default_styles(self):
        style = self.doc.styles['Normal']
        font = style.font
        font.name = self.font_name
        font.size = Pt(12)
        font.color.rgb = BLACK

    def _clean_text(self, text):
        if not text:
            return ""
        # R-02 Hard Gate: Larang em-dash, ubah ke simbol standar
        return text.replace("—", "-")

    def add_cover(self, title, subtitle=None, author=None, nim=None, major=None, institution=None, year=None):
        """Membuat Halaman Judul / Cover Formal."""
        title = self._clean_text(title)
        
        # Jarak atas cover
        p_top = self.doc.add_paragraph()
        p_top.paragraph_format.space_before = Pt(36)
        
        # Judul Utama
        p_title = self.doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(8)
        r_title = p_title.add_run(title.upper())
        r_title.font.name = self.font_name
        r_title.font.size = Pt(15)
        r_title.font.bold = True
        r_title.font.color.rgb = BLACK

        # Subjudul (jika ada)
        if subtitle:
            subtitle = self._clean_text(subtitle)
            p_sub = self.doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(24)
            r_sub = p_sub.add_run(subtitle)
            r_sub.font.name = self.font_name
            r_sub.font.size = Pt(12)
            r_sub.font.italic = True
            r_sub.font.color.rgb = BLACK

        # Ruang pemisah / placeholder logo
        p_mid = self.doc.add_paragraph()
        p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_mid.paragraph_format.space_before = Pt(72)
        p_mid.paragraph_format.space_after = Pt(72)
        r_mid = p_mid.add_run("[ LOGO INSTANSI / UNIVERSITAS ]")
        r_mid.font.name = self.font_name
        r_mid.font.size = Pt(10)
        r_mid.font.italic = True
        r_mid.font.color.rgb = BLACK

        # Identitas Penulis
        p_auth = self.doc.add_paragraph()
        p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_auth.paragraph_format.space_after = Pt(2)
        r_a1 = p_auth.add_run("Disusun Oleh:\n")
        r_a1.font.name = self.font_name
        r_a1.font.size = Pt(11)
        r_a1.font.color.rgb = BLACK
        
        if author:
            r_a2 = p_auth.add_run(f"{author}\n")
            r_a2.font.name = self.font_name
            r_a2.font.size = Pt(12)
            r_a2.font.bold = True
            r_a2.font.color.rgb = BLACK
        if nim:
            r_a3 = p_auth.add_run(f"NIM: {nim}\n")
            r_a3.font.name = self.font_name
            r_a3.font.size = Pt(11)
            r_a3.font.color.rgb = BLACK

        # Footer Instansi & Tahun
        p_bot = self.doc.add_paragraph()
        p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_bot.paragraph_format.space_before = Pt(80)
        p_bot.paragraph_format.space_after = Pt(0)
        
        info_lines = []
        if major: info_lines.append(major.upper())
        if institution: info_lines.append(institution.upper())
        if year: info_lines.append(str(year))
        
        r_bot = p_bot.add_run("\n".join(info_lines))
        r_bot.font.name = self.font_name
        r_bot.font.size = Pt(12)
        r_bot.font.bold = True
        r_bot.font.color.rgb = BLACK

        # Pindah halaman setelah cover
        self.doc.add_page_break()

    def add_heading_1(self, text):
        text = self._clean_text(text)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text.upper())
        run.font.name = self.font_name
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = BLACK
        return p

    def add_heading_2(self, text):
        text = self._clean_text(text)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = self.font_name
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = BLACK
        return p

    def add_heading_3(self, text):
        text = self._clean_text(text)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = self.font_name
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = BLACK
        return p

    def add_paragraph(self, text, space_after=6, line_spacing=1.15):
        text = self._clean_text(text)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        run = p.add_run(text)
        run.font.name = self.font_name
        run.font.size = Pt(11.5)
        run.font.color.rgb = BLACK
        return p

    def add_bullet(self, bold_prefix, text):
        bold_prefix = self._clean_text(bold_prefix)
        text = self._clean_text(text)
        p = self.doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        
        r1 = p.add_run(bold_prefix)
        r1.font.name = self.font_name
        r1.font.size = Pt(11.5)
        r1.font.bold = True
        r1.font.color.rgb = BLACK
        
        r2 = p.add_run(text)
        r2.font.name = self.font_name
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = BLACK
        return p

    def add_table(self, headers, rows_data):
        """Membuat tabel bergaris rapi dengan header abu-abu dan teks hitam."""
        table = self.doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Border tabel rapi
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="888888"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="888888"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(borders)

        # Header Row
        for col_idx, h in enumerate(headers):
            cell = table.cell(0, col_idx)
            # Shading abu-abu muda
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E0E0E0"/>')
            tcPr.append(shd)
            # Padding sel
            tcMar = parse_xml(f'''
                <w:tcMar {nsdecls("w")}>
                    <w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/>
                    <w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/>
                </w:tcMar>
            ''')
            tcPr.append(tcMar)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(self._clean_text(h))
            r.font.name = self.font_name
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = BLACK

        # Data Rows
        for row_idx, row in enumerate(rows_data, start=1):
            bg_col = "F8F8F8" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(row):
                cell = table.cell(row_idx, col_idx)
                tcPr = cell._tc.get_or_add_tcPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_col}"/>')
                tcPr.append(shd)
                tcMar = parse_xml(f'''
                    <w:tcMar {nsdecls("w")}>
                        <w:top w:w="60" w:type="dxa"/><w:bottom w:w="60" w:type="dxa"/>
                        <w:left w:w="90" w:type="dxa"/><w:right w:w="90" w:type="dxa"/>
                    </w:tcMar>
                ''')
                tcPr.append(tcMar)
                
                p = cell.paragraphs[0]
                r = p.add_run(self._clean_text(str(val)))
                r.font.name = self.font_name
                r.font.size = Pt(9.5)
                r.font.color.rgb = BLACK
                if col_idx == 0:
                    r.font.bold = True

        p_after = self.doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(4)
        p_after.paragraph_format.space_after = Pt(6)

    def add_callout(self, text, bold_title="Catatan: "):
        """Menambahkan callout box bergaris aksen kiri."""
        text = self._clean_text(text)
        bold_title = self._clean_text(bold_title)
        
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.3)
        
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>')
        tcPr.append(shd)
        
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:left w:val="single" w:sz="18" w:space="0" w:color="333333"/>
                <w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)

        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="100" w:type="dxa"/><w:bottom w:w="100" w:type="dxa"/>
                <w:left w:w="150" w:type="dxa"/><w:right w:w="150" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r1 = p.add_run(bold_title)
        r1.font.name = self.font_name
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = BLACK
        
        r2 = p.add_run(text)
        r2.font.name = self.font_name
        r2.font.size = Pt(11)
        r2.font.color.rgb = BLACK

        p_after = self.doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(2)
        p_after.paragraph_format.space_after = Pt(6)

    def add_figure(self, image_path, caption, width_inches=6.0):
        """Menyisipkan gambar diagram dengan caption terpusat."""
        if not os.path.exists(image_path):
            return
        p_img = self.doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        self.doc.add_picture(image_path, width=Inches(width_inches))

        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        r = p_cap.add_run(self._clean_text(caption))
        r.font.name = self.font_name
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = BLACK

    def save(self, output_path, open_explorer=True):
        """Menyimpan dokumen dan membuka explorer."""
        self.doc.save(output_path)
        print(f"[StandardDocument] File berhasil disimpan di: {output_path}")
        if open_explorer and sys.platform == "win32":
            try:
                subprocess.Popen(f'explorer.exe /select,"{os.path.abspath(output_path)}"')
            except Exception:
                pass
        return output_path
